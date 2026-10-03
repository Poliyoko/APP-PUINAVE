"""Importador read-only de la fuente léxica certificada REAL-508.

Preserva exactamente la grafía lingüística del Excel.
No modifica Excel, audio ni almacenamiento productivo.
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from sgoda.integration.real508_adapter import (
    Real508StorageConfig,
    canonical_real508_id,
)
from sgoda.integration.real508_lexical_integration import (
    build_real508_lexical_entry,
    load_real508_into_operational_repository,
)
from sgoda.lexical_engine.models import LexicalEntry
from sgoda.operational_platform.database import OperationalRepository


SHEET_NAME = "hoja1"

COLUMN_ID = "ID"
COLUMN_PUINAVE = "PALABRA EN PUINAVE"
COLUMN_WRITING_PUINAVE = "ESCRITURA EN PUINAVE"
COLUMN_SPANISH = "PALABRA EN ESPAÑOL"
COLUMN_AUDIO_LOT = "LOTTE DE AUDIO EN PUINAVE"

REQUIRED_COLUMNS = (
    COLUMN_ID,
    COLUMN_PUINAVE,
    COLUMN_WRITING_PUINAVE,
    COLUMN_SPANISH,
    COLUMN_AUDIO_LOT,
)


def _exact_text(value: object) -> str:
    """Convierte a texto sin strip, normalización ni reinterpretación."""
    if value is None:
        return ""
    return str(value)


def _source_id(value: object) -> str:
    if value is None:
        raise ValueError("ID vacío en fuente léxica.")

    if isinstance(value, int):
        return canonical_real508_id(value)

    if isinstance(value, float) and value.is_integer():
        return canonical_real508_id(int(value))

    return canonical_real508_id(str(value))


def iter_certified_real508_entries(
    excel_path: str | Path,
    storage_config: Real508StorageConfig,
) -> Iterable[LexicalEntry]:
    """Lee hoja1 y produce LexicalEntry preservando texto exacto."""

    path = Path(excel_path)

    if not path.is_file():
        raise FileNotFoundError(path)

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=False,
    )

    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError(
                f"Hoja canónica ausente: {SHEET_NAME}"
            )

        worksheet = workbook[SHEET_NAME]
        rows = worksheet.iter_rows(values_only=True)

        try:
            header_row = next(rows)
        except StopIteration as exc:
            raise ValueError("Excel léxico vacío.") from exc

        headers = tuple(
            "" if value is None else str(value)
            for value in header_row
        )

        missing = [
            name
            for name in REQUIRED_COLUMNS
            if name not in headers
        ]

        if missing:
            raise ValueError(
                "Columnas canónicas faltantes: "
                + ", ".join(missing)
            )

        indexes = {
            name: headers.index(name)
            for name in REQUIRED_COLUMNS
        }

        count = 0

        for excel_row, row in enumerate(rows, start=2):
            values = list(row)

            if len(values) < len(headers):
                values.extend(
                    [None] * (len(headers) - len(values))
                )

            selected = {
                name: values[indexes[name]]
                for name in REQUIRED_COLUMNS
            }

            if all(
                value is None or str(value) == ""
                for value in selected.values()
            ):
                continue

            entry_id = _source_id(
                selected[COLUMN_ID]
            )

            expected_id = canonical_real508_id(
                count + 1
            )

            if entry_id != expected_id:
                raise ValueError(
                    "Secuencia REAL-508 inválida en fila "
                    f"{excel_row}: esperado={expected_id}, "
                    f"actual={entry_id}"
                )

            puinave = _exact_text(
                selected[COLUMN_PUINAVE]
            )

            writing_puinave = _exact_text(
                selected[COLUMN_WRITING_PUINAVE]
            )

            spanish = _exact_text(
                selected[COLUMN_SPANISH]
            )

            audio_lot = _exact_text(
                selected[COLUMN_AUDIO_LOT]
            )

            if not puinave:
                raise ValueError(
                    f"Puinave vacío: {entry_id}"
                )

            if not writing_puinave:
                raise ValueError(
                    f"Escritura Puinave vacía: {entry_id}"
                )

            if not spanish:
                raise ValueError(
                    f"Español vacío: {entry_id}"
                )

            if not audio_lot:
                raise ValueError(
                    f"Lote de audio vacío: {entry_id}"
                )

            metadata = {
                "fuente_lexica": "EXCEL_CERTIFICADO",
                "hoja_fuente": SHEET_NAME,
                "fila_excel": excel_row,
                "escritura_puinave": writing_puinave,
                "lote_audio_puinave": audio_lot,
                "grafia_puinave_autoritativa": True,
                "preservar_texto_exacto": True,
                "normalizar_puinave": False,
                "corregir_automaticamente_puinave": False,
                "reinterpretar_puinave": False,
            }


            entry = build_real508_lexical_entry(
                lexical_id=entry_id,
                config=storage_config,
                puinave=puinave,
                spanish=spanish,
                extra=metadata,
            )

            count += 1
            yield entry

        if count != 508:
            raise ValueError(
                f"REAL-508 requiere 508 registros; actual={count}"
            )

    finally:
        workbook.close()


def load_certified_real508_excel(
    excel_path: str | Path,
    storage_config: Real508StorageConfig,
    repository: OperationalRepository | None = None,
) -> OperationalRepository:
    """Carga los 508 registros en el repositorio operativo en memoria."""

    target = repository or OperationalRepository()

    entries = tuple(
        iter_certified_real508_entries(
            excel_path,
            storage_config,
        )
    )

    if len(entries) != 508:
        raise ValueError(
            f"REAL-508 requiere 508 entradas; actual={len(entries)}"
        )

    load_real508_into_operational_repository(
        repository=target,
        entries=entries,
    )

    if target.count() != 508:
        raise ValueError(
            "Repositorio operativo no contiene "
            f"508 registros: actual={target.count()}"
        )

    return target
